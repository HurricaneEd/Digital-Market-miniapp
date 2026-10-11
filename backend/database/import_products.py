"""将清洗后的商品 Excel 或采集 JSON 导入商品目录。

默认工作流是：

1. 爬虫生成 products_raw.xlsx / products_raw.json；
2. 人工在 Excel 的 cleaning_result 工作表中清洗数据；
3. 将 approved 行导入 catalog_products。

示例：
    python backend/database/import_products.py data/raw/20260903_195829/products_raw.xlsx
    python backend/database/import_products.py data/raw/20260903_195829/products_raw.json --allow-unapproved
    python backend/database/import_products.py data/raw/20260903_195829/products_raw.xlsx --dry-run

脚本不会导入评论，也不会处理 source_url、params_url、sub_images 等字段。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

DEFAULT_SETTINGS_MODULE = 'DigRecSys.settings'
CATEGORY_NAMES = {
    'digital_camera': '数码相机',
    'cell_phone': '智能手机',
    'mice': '鼠标',
    'microphone': '耳机',
    'pocketpower': '移动电源',
    'tablepc': '平板电脑',
}
APPROVED_STATUS = 'approved'
EMPTY_PRICE_VALUES = {'', '-', '--', '暂无报价', '面议', '无报价', 'n/a', 'na', 'null', 'none'}
TRUE_VALUES = {'1', 'true', 'yes', 'y', '是', '在售', 'active'}
FALSE_VALUES = {'0', 'false', 'no', 'n', '否', '下架', 'inactive'}


class ImportValidationError(ValueError):
    """输入文件不符合商品导入契约。"""


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any, field_name: str, row_number: int, required: bool = False) -> str:
    if _is_blank(value):
        if required:
            raise ImportValidationError(f'第 {row_number} 行缺少 {field_name}')
        return ''
    result = str(value).strip()
    if required and not result:
        raise ImportValidationError(f'第 {row_number} 行缺少 {field_name}')
    return result


def _first_value(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row and not _is_blank(row[name]):
            return row[name]
    return None


def _parse_price(value: Any, row_number: int) -> Decimal:
    if _is_blank(value):
        return Decimal('0.00')
    if isinstance(value, Decimal):
        parsed = value
    else:
        raw = str(value).strip()
        if raw.lower() in EMPTY_PRICE_VALUES:
            return Decimal('0.00')
        normalized = raw.replace(',', '').replace('￥', '').replace('¥', '').strip()
        match = re.search(r'\d+(?:\.\d+)?', normalized)
        if not match:
            raise ImportValidationError(f'第 {row_number} 行价格无法解析：{raw}')
        try:
            parsed = Decimal(match.group(0))
        except InvalidOperation as exc:
            raise ImportValidationError(f'第 {row_number} 行价格无法解析：{raw}') from exc

    if parsed < 0:
        raise ImportValidationError(f'第 {row_number} 行价格不能为负数')
    parsed = parsed.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    if parsed > Decimal('99999999.99'):
        raise ImportValidationError(f'第 {row_number} 行价格超过数据库范围：{parsed}')
    return parsed


def _parse_stock(value: Any, row_number: int) -> int:
    if _is_blank(value):
        return 0
    try:
        parsed = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ImportValidationError(f'第 {row_number} 行库存无法解析：{value}') from exc
    if parsed < 0 or parsed != parsed.to_integral_value():
        raise ImportValidationError(f'第 {row_number} 行库存必须是非负整数：{value}')
    result = int(parsed)
    if result > 2147483647:
        raise ImportValidationError(f'第 {row_number} 行库存超过数据库范围：{result}')
    return result


def _parse_bool(value: Any, row_number: int) -> bool:
    if _is_blank(value):
        return True
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    normalized = str(value).strip().lower()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    raise ImportValidationError(f'第 {row_number} 行 is_active 无法解析：{value}')


def _parse_all_params(value: Any, row_number: int) -> dict[str, Any]:
    if _is_blank(value):
        return {}
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError as exc:
            raise ImportValidationError(f'第 {row_number} 行 all_params 不是有效 JSON：{exc.msg}') from exc
    if not isinstance(value, dict):
        raise ImportValidationError(f'第 {row_number} 行 all_params 顶层必须是 JSON 对象')
    invalid_groups = [group for group, params in value.items() if not isinstance(group, str) or not isinstance(params, dict)]
    if invalid_groups:
        raise ImportValidationError(f'第 {row_number} 行 all_params 的参数分组必须是 JSON 对象')
    return value


def _parse_img_url(value: Any, row_number: int) -> str:
    image_url = _text(value, 'img_url', row_number)
    if image_url.startswith('//'):
        image_url = f'https:{image_url}'
    if len(image_url) > 1000:
        raise ImportValidationError(f'第 {row_number} 行 img_url 超过 1000 个字符')
    return image_url


def normalize_row(row: dict[str, Any], row_number: int) -> dict[str, Any]:
    """把 Excel/JSON 的展示值转换为 Product 模型可接受的值。"""
    prod_id = _text(row.get('prod_id'), 'prod_id', row_number, required=True)
    if len(prod_id) > 50:
        raise ImportValidationError(f'第 {row_number} 行 prod_id 超过 50 个字符')

    category_code = _text(row.get('category_code'), 'category_code', row_number, required=True)
    if len(category_code) > 50:
        raise ImportValidationError(f'第 {row_number} 行 category_code 超过 50 个字符')

    category_name = _text(
        _first_value(row, 'category_name', 'category_name_clean', 'category_name_raw'),
        'category_name',
        row_number,
    ) or CATEGORY_NAMES.get(category_code, category_code)
    brand = _text(_first_value(row, 'brand', 'brand_clean', 'brand_raw'), 'brand', row_number) or '其他'
    name = _text(_first_value(row, 'name', 'name_clean', 'name_raw'), 'name', row_number, required=True)
    if len(brand) > 50:
        raise ImportValidationError(f'第 {row_number} 行 brand 超过 50 个字符')
    if len(category_name) > 50:
        raise ImportValidationError(f'第 {row_number} 行 category_name 超过 50 个字符')
    if len(name) > 200:
        raise ImportValidationError(f'第 {row_number} 行 name 超过 200 个字符')

    return {
        'prod_id': prod_id,
        'category_code': category_code,
        'category_name': category_name,
        'brand': brand,
        'name': name,
        'price': _parse_price(_first_value(row, 'price', 'price_clean', 'price_raw'), row_number),
        'img_url': _parse_img_url(row.get('img_url') or row.get('image_url_raw'), row_number),
        'all_params': _parse_all_params(row.get('all_params'), row_number),
        'stock': _parse_stock(row.get('stock'), row_number),
        'is_active': _parse_bool(row.get('is_active'), row_number),
    }


def read_json_rows(path: Path) -> list[tuple[int, dict[str, Any]]]:
    with path.open('r', encoding='utf-8-sig') as file:
        payload = json.load(file)
    if isinstance(payload, dict):
        payload = payload.get('products', payload.get('data'))
    if not isinstance(payload, list):
        raise ImportValidationError('JSON 根节点必须是商品数组，或包含 products/data 数组')
    rows = []
    for row_number, row in enumerate(payload, start=1):
        if not isinstance(row, dict):
            raise ImportValidationError(f'第 {row_number} 行 JSON 必须是对象')
        rows.append((row_number, row))
    return rows


def read_excel_rows(path: Path, sheet_name: str) -> list[tuple[int, dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError('读取 Excel 需要安装 openpyxl：pip install openpyxl') from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ImportValidationError(
                f'Excel 不存在工作表 {sheet_name}，可选工作表：{", ".join(workbook.sheetnames)}'
            )
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration as exc:
            raise ImportValidationError(f'Excel 工作表 {sheet_name} 为空') from exc

        headers = [str(value).strip() if value is not None else '' for value in header_values]
        if not headers or not headers[0]:
            raise ImportValidationError(f'Excel 工作表 {sheet_name} 缺少表头')
        if len(headers) != len(set(headers)):
            raise ImportValidationError(f'Excel 工作表 {sheet_name} 存在重复表头')

        result = []
        for row_number, values in enumerate(rows, start=2):
            row = dict(zip(headers, values))
            if all(_is_blank(value) for value in row.values()):
                continue
            result.append((row_number, row))
        return result
    finally:
        workbook.close()


def read_rows(path: Path, sheet_name: str) -> tuple[str, list[tuple[int, dict[str, Any]]]]:
    suffix = path.suffix.lower()
    if suffix == '.json':
        return 'json', read_json_rows(path)
    if suffix in {'.xlsx', '.xlsm'}:
        return 'excel', read_excel_rows(path, sheet_name)
    raise ImportValidationError('输入文件必须是 .json、.xlsx 或 .xlsm')


def _approved_rows(
    rows: Iterable[tuple[int, dict[str, Any]]],
    allow_unapproved: bool,
) -> tuple[list[tuple[int, dict[str, Any]]], int]:
    selected = []
    skipped = 0
    for row_number, row in rows:
        status = _text(row.get('status'), 'status', row_number).lower()
        if status and status != APPROVED_STATUS and not allow_unapproved:
            skipped += 1
            continue
        if not status and not allow_unapproved:
            skipped += 1
            continue
        selected.append((row_number, row))
    return selected, skipped


def _normalize_rows(rows: Iterable[tuple[int, dict[str, Any]]]) -> list[dict[str, Any]]:
    normalized = []
    seen: dict[str, int] = {}
    errors = []
    for row_number, row in rows:
        try:
            product = normalize_row(row, row_number)
            if product['prod_id'] in seen:
                raise ImportValidationError(
                    f'prod_id {product["prod_id"]} 与第 {seen[product["prod_id"]]} 行重复'
                )
            seen[product['prod_id']] = row_number
            normalized.append(product)
        except ImportValidationError as exc:
            errors.append(str(exc))
    if errors:
        details = '\n'.join(f'  - {error}' for error in errors)
        raise ImportValidationError(f'共有 {len(errors)} 行无法导入：\n{details}')
    return normalized


def import_to_database(products: list[dict[str, Any]], dry_run: bool = False) -> dict[str, int]:
    """在一个事务中完成品牌、分类和商品的幂等导入。"""
    from django.apps import apps

    if not apps.ready:
        import django

        django.setup()
    from django.db import transaction
    from Product.models import Brand, Category, Product

    result = {'created': 0, 'updated': 0, 'brands_created': 0, 'categories_created': 0}
    with transaction.atomic():
        for item in products:
            brand, brand_created = Brand.objects.get_or_create(name=item['brand'])
            if brand_created:
                result['brands_created'] += 1

            category, category_created = Category.objects.get_or_create(
                code=item['category_code'],
                defaults={'name': item['category_name']},
            )
            if category_created:
                result['categories_created'] += 1

            _, product_created = Product.objects.update_or_create(
                prod_id=item['prod_id'],
                defaults={
                    'brand': brand,
                    'category': category,
                    'name': item['name'],
                    'price': item['price'],
                    'img_url': item['img_url'],
                    'all_params': item['all_params'],
                    'stock': item['stock'],
                    'is_active': item['is_active'],
                },
            )
            result['created' if product_created else 'updated'] += 1

        if dry_run:
            transaction.set_rollback(True)
    return result


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description='导入商品清洗结果或采集 JSON 到 Django 商品目录')
    parser.add_argument('input_file', type=Path, help='products_raw.xlsx 或 products_raw.json')
    parser.add_argument(
        '--sheet',
        default='cleaning_result',
        help='Excel 工作表，默认 cleaning_result；读取 products_raw 时需显式指定 --sheet products_raw',
    )
    parser.add_argument(
        '--allow-unapproved',
        action='store_true',
        help='允许导入 pending/无 status 的原始数据；默认只导入 status=approved 的行',
    )
    parser.add_argument('--dry-run', action='store_true', help='执行完整校验和事务写入流程，但回滚所有数据库变更')
    parser.add_argument('--settings', default=DEFAULT_SETTINGS_MODULE, help='Django settings 模块')
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = _build_parser()
    args = parser.parse_args(argv)
    input_file = args.input_file.resolve()
    if not input_file.is_file():
        parser.error(f'输入文件不存在：{input_file}')

    try:
        source_type, rows = read_rows(input_file, args.sheet)
        selected_rows, skipped = _approved_rows(rows, args.allow_unapproved)
        if not selected_rows:
            raise ImportValidationError(
                '没有可导入记录；请在 cleaning_result 将 status 标记为 approved，'
                '或对原始 JSON/工作表使用 --allow-unapproved。'
            )
        products = _normalize_rows(selected_rows)

        os.environ.setdefault('DJANGO_SETTINGS_MODULE', args.settings)
        import django

        django.setup()
        result = import_to_database(products, dry_run=args.dry_run)
    except (ImportValidationError, RuntimeError, OSError, json.JSONDecodeError) as exc:
        print(f'[失败] {exc}', file=sys.stderr)
        return 1
    except Exception as exc:
        print(f'[数据库失败] {exc}', file=sys.stderr)
        print('请先执行 python backend/manage.py migrate，再重新运行导入脚本。', file=sys.stderr)
        return 1

    action = '预检完成（已回滚）' if args.dry_run else '入库完成'
    print(
        f'[{action}] 来源={source_type}，读取={len(rows)}，跳过={skipped}，'
        f'商品新增={result["created"]}，更新={result["updated"]}，'
        f'品牌新增={result["brands_created"]}，分类新增={result["categories_created"]}'
    )
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
